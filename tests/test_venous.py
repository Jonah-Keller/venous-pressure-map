"""Tests for the app's logic. Run: pytest -q"""
import base64
import io
import json

import numpy as np
from openpyxl import load_workbook

from venous import anatomy, demo, export, sinus3d
from venous.anatomy import Measurement as M


def test_units_round_trip():
    assert anatomy.convert(20, "mmHg", "cm H₂O") == 27.2
    assert anatomy.convert(27.2, "cm H₂O", "mmHg") == 20.0


def test_narrowing_is_found_from_the_largest_drop():
    ms = demo.example_measurements("mmHg")
    s = anatomy.find_stenosis(ms, "mmHg")
    assert (s.up, s.down, s.drop) == ("Transverse sinus (lateral)", "Transverse-sigmoid junction", 17)
    after = anatomy.find_stenosis(ms, "mmHg", after=True)
    assert (after.up, after.down, after.drop) == (s.up, s.down, 1)   # same place, after values


def test_narrowing_moves_with_the_pressures():
    ms = [M("Torcula", "", 30), M("Transverse sinus (mid)", "Left", 29),
          M("Transverse sinus (lateral)", "Left", 28), M("Sigmoid sinus", "Left", 27),
          M("Jugular bulb", "Left", 10)]
    s = anatomy.find_stenosis(ms, "mmHg")
    assert (s.up, s.down) == ("Sigmoid sinus", "Jugular bulb")


def test_no_narrowing_when_drops_are_small_and_straight_sinus_ignored():
    ms = [M("Superior sagittal sinus", "", 12), M("Straight sinus", "", 30),
          M("Torcula", "", 11), M("Sigmoid sinus", "Right", 10)]
    assert anatomy.find_stenosis(ms, "mmHg") is None


def test_record_and_exports():
    ms = demo.example_measurements("cm H₂O")
    row = export.record(ms, "cm H₂O", "Right", "VP-0001", "2026-10-22")
    assert row["study_id"] == "VP-0001" and row["unit"] == "cm H2O"
    assert row["transverse_sinus_lateral_pre"] == 39 and row["gradient_pre"] == 23
    assert "mrn" not in " ".join(row).lower()
    assert b"VP-0001" in export.csv_bytes(row)
    wb = load_workbook(io.BytesIO(export.excel_bytes(ms, "cm H₂O", row)))
    assert wb.sheetnames == ["Readings", "Record", "Long format"]
    long = export.long_format(ms, "cm H₂O")
    assert set(long.phase) == {"pre_stent", "post_stent"} and {"midline", "right"} == set(long.side)
    assert {"medial", "mid", "lateral"} <= set(long.sub_location.dropna())


def test_viewer_payload():
    ms = demo.example_measurements("mmHg")
    d = sinus3d.scan()
    dec = lambda b, t: np.frombuffer(base64.b64decode(b), t)              # noqa: E731
    pay = sinus3d.viewer_payload(ms, "mmHg", "Right")
    col = dec(pay["mesh"]["colour_before"], np.uint8).reshape(-1, 3)
    grey = (col == sinus3d.GREY).all(1)
    left_ts = (d["v_side"] == 0) & (d["v_idx"] > d["lm_left"][0] + 3)
    assert grey[left_ts].mean() > 0.9                                       # left not measured
    vb = dec(pay["mesh"]["verts_before"], np.float32).reshape(-1, 3)
    moved = np.linalg.norm(vb - d["verts"], axis=1)
    assert moved.max() > 1 and moved[d["v_side"] != 1].max() == 0         # only right narrowed
    assert pay["drop_before"] == "17 mmHg" and pay["drop_after"] == "1 mmHg"
    labels = pay["labels_before"]
    assert all(lab["value"].isdigit() for lab in labels)                    # whole numbers only
    assert all("·" not in lab["value"] for lab in labels)
    flow = {p["side"]: p for p in pay["paths"]}
    assert flow["Right"]["q_before"] < flow["Right"]["q_after"]
    # narrowing follows the pressures: a drop lower down moves it into the sigmoid
    low = [M("Torcula", "", 30), M("Transverse sinus (lateral)", "Right", 29),
           M("Sigmoid sinus", "Right", 28), M("Jugular bulb", "Right", 10)]
    assert sinus3d.viewer_payload(low, "mmHg", "Right")["stenosis"]["pos"] != pay["stenosis"]["pos"]
    html = sinus3d.viewer_html(pay)
    assert "__PAYLOAD__" not in html and json.loads(json.dumps(pay))
    assert sinus3d.viewer_payload([M("Torcula", "", 30)], "mmHg", "Left")["stenosis"] is None
    assert "TopBrain" in sinus3d.credit()
